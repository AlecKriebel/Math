"""Verify the narrow edit and preserve font metadata omitted on export.

Workbook content is authored by Artifact Tool. openpyxl is read-only here.
The ZIP repair restores an unsupported round-trip metadata field only.
"""
import copy
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl

base = Path(__file__).resolve().parent
source = Path('/Users/alec/Downloads/Papers (4).xlsx')
output = base / 'outputs/01a09d58-c65d-7731-a4d1-086aef7a9db0/Papers - SimAmpA assessment.xlsx'
a = openpyxl.load_workbook(source)
b = openpyxl.load_workbook(output)
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with ZipFile(output) as archive:
    entries = [(i, archive.read(i.filename)) for i in archive.infolist()]
tree = ET.fromstring(dict((i.filename, v) for i,v in entries)['xl/styles.xml'])
fonts = tree.find('s:fonts', ns)
assert len(a._fonts) == len(b._fonts)
for sheet in a:
    for row in sheet:
        for cell in row:
            assert cell._style is None or cell._style.fontId == b[sheet.title][cell.coordinate]._style.fontId
with ZipFile(source) as archive:
    original_fonts = ET.fromstring(archive.read('xl/styles.xml')).find('s:fonts', ns)
restored = sum(x != y for x,y in zip(a._fonts,b._fonts))
if restored:
    index = list(tree).index(fonts)
    tree.remove(fonts)
    tree.insert(index, original_fonts)
    replacement = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
    tmp = output.with_suffix('.tmp')
    with ZipFile(tmp, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, replacement if info.filename == 'xl/styles.xml' else data)
    tmp.replace(output)

b = openpyxl.load_workbook(output)
cached = openpyxl.load_workbook(output, data_only=True)
expected = json.loads((base/'workbook_updates.json').read_text())
assert a.sheetnames == b.sheetnames
value_changes, style_changes = [], []
for sheet in a:
    result = b[sheet.title]
    assert sheet.freeze_panes == result.freeze_panes
    assert str(sheet.merged_cells) == str(result.merged_cells)
    assert str(sheet.data_validations) == str(result.data_validations)
    assert sheet.auto_filter == result.auto_filter
    assert len(sheet.conditional_formatting) == len(result.conditional_formatting)
    for key in sheet.conditional_formatting:
        assert sheet.conditional_formatting[key] == result.conditional_formatting[key]
    assert len(sheet._charts) == len(result._charts) == 0
    assert len(sheet._images) == len(result._images) == 0
    for row in sheet:
        for cell in row:
            other = result[cell.coordinate]
            if cell.value != other.value or cell.data_type != other.data_type:
                assert sheet.title == 'Papers' and cell.coordinate in expected
                assert other.value == expected[cell.coordinate]
                value_changes.append(cell.coordinate)
            for attribute in ['font','fill','border','alignment','number_format','protection']:
                if copy.copy(getattr(cell,attribute)) != copy.copy(getattr(other,attribute)):
                    assert sheet.title == 'Papers' and cell.coordinate in ['E12','F12','K12'] and attribute == 'alignment', (sheet.title,cell.coordinate,attribute)
                    style_changes.append([cell.coordinate,attribute])
            assert cell.comment == other.comment
            if cell.hyperlink or other.hyperlink:
                assert (cell.hyperlink.target if cell.hyperlink else None) == (other.hyperlink.target if other.hyperlink else None)
assert sorted(value_changes) == sorted(expected)
assert b['Papers']['L12'].value == '=AVERAGE(M12:N12)'
assert cached['Papers']['L12'].value == 7.25
assert b['Papers']['N12'].value == 7.2
assert b['Papers'].row_dimensions[12].height == 200
report = {'modified_cells': value_changes, 'alignment_changes':style_changes,
          'average_formula_preserved':True,'average_cached_value':7.25,
          'original_font_records_restored':restored,
          'all_other_values_types_formulas_cell_styles_and_checked_features_preserved':True}
(base/'workbook_verification.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
