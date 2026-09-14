"""Preserve unrelated OOXML after Artifact Tool authors the requested cells.

The exporter drops original font-scheme metadata. Transplant its authored
values and local alignments into the original package instead. All other
worksheet/object parts remain byte-identical. No spreadsheet authoring library
is used here; Artifact Tool remains the source of the new values and layout.
"""
from pathlib import Path
from copy import deepcopy, copy
import json
import zipfile
import xml.etree.ElementTree as ET
import openpyxl  # read-only verification

ROOT = Path(__file__).resolve().parent
SOURCE = Path('/Users/alec/Downloads/Papers (4).xlsx')
OUTPUT = ROOT / 'outputs/turing-review/Papers - Turing next steps.xlsx'
N = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ET.register_namespace('', N[1:-1])
ET.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
with zipfile.ZipFile(SOURCE) as archive:
    original = {name: archive.read(name) for name in archive.namelist()}
with zipfile.ZipFile(OUTPUT) as archive:
    authored = {name: archive.read(name) for name in archive.namelist()}
sheet_part = 'xl/worksheets/sheet1.xml'
source_sheet = ET.fromstring(original[sheet_part])
output_sheet = ET.fromstring(authored[sheet_part])
source_styles = ET.fromstring(original['xl/styles.xml'])
output_styles = ET.fromstring(authored['xl/styles.xml'])
source_xfs = source_styles.find(N+'cellXfs')
output_xfs = output_styles.find(N+'cellXfs')
strings = ET.fromstring(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else []
source_cells = {c.attrib['r']: c for c in source_sheet.iter(N+'c')}
output_cells = {c.attrib['r']: c for c in output_sheet.iter(N+'c')}
rows = {r.attrib['r']: r for r in source_sheet.find(N+'sheetData')}
changed = ['X1','C8','E8','F8','M8','P8','W8','X8']
for address in changed:
    old = source_cells.get(address)
    new = deepcopy(output_cells[address])
    # Expand only newly authored shared strings, keeping original sst untouched.
    if new.get('t') == 's':
        item = strings[int(new.find(N+'v').text)]
        new.remove(new.find(N+'v'))
        new.set('t','inlineStr')
        inline = ET.SubElement(new,N+'is')
        for child in item:
            inline.append(deepcopy(child))
    # Preserve original font/fill/border/number format, using new alignment only.
    style_source = source_cells.get('W1') if address == 'X1' else old
    xf = deepcopy(source_xfs[int(style_source.get('s','0')) if style_source is not None else 0])
    authored_xf = output_xfs[int(new.get('s','0'))]
    align = authored_xf.find(N+'alignment')
    existing_align = xf.find(N+'alignment')
    if existing_align is not None:
        xf.remove(existing_align)
    if align is not None:
        xf.append(deepcopy(align))
        xf.set('applyAlignment','1')
    new.set('s',str(len(source_xfs)))
    source_xfs.append(xf)
    row = rows[''.join(x for x in address if x.isdigit())]
    if old is not None:
        row.remove(old)
    row.append(new)
source_xfs.set('count',str(len(source_xfs)))
# Preserve the original average formula and refresh only its cached value.
value = source_cells['L8'].find(N+'v')
if value is None:
    value = ET.SubElement(source_cells['L8'],N+'v')
value.text = output_cells['L8'].find(N+'v').text
rows['8'].set('ht','210')
rows['8'].set('customHeight','1')
for row in [rows['1'],rows['8']]:
    cells = list(row)
    def col_index(cell):
        result=0
        for c in cell.attrib['r']:
            if c.isalpha(): result=result*26+ord(c)-64
        return result
    row[:] = sorted(cells,key=col_index)
cols = source_sheet.find(N+'cols')
# Width edits are limited to the newly populated notes/source columns.
for index, width in [(23,'55'),(24,'78')]:
    for col in list(cols):
        lo,hi=int(col.get('min')),int(col.get('max'))
        if lo<=index<=hi:
            cols.remove(col)
            if lo<index:
                left=deepcopy(col);left.set('max',str(index-1));cols.append(left)
            if hi>index:
                right=deepcopy(col);right.set('min',str(index+1));cols.append(right)
    ET.SubElement(cols,N+'col',{'min':str(index),'max':str(index),'width':width,'customWidth':'1'})
cols[:]=sorted(list(cols),key=lambda c:int(c.get('min')))
original[sheet_part]=ET.tostring(source_sheet,encoding='utf-8',xml_declaration=True)
original['xl/styles.xml']=ET.tostring(source_styles,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(OUTPUT,'w',zipfile.ZIP_DEFLATED) as archive:
    for name,data in original.items(): archive.writestr(name,data)

# Read-only semantic verification, including all unrelated populated cells.
a=openpyxl.load_workbook(SOURCE)
b=openpyxl.load_workbook(OUTPUT)
assert a.sheetnames==b.sheetnames
value_changes=[]
style_changes=[]
for sheet in a:
    target=b[sheet.title]
    for row in sheet:
        for cell in row:
            other=target[cell.coordinate]
            if cell.value!=other.value:
                assert sheet.title=='Papers' and cell.coordinate in changed
                value_changes.append(cell.coordinate)
            for key in ['font','fill','border','alignment','number_format','protection']:
                if copy(getattr(cell,key))!=copy(getattr(other,key)):
                    assert sheet.title=='Papers' and cell.coordinate in changed
                    style_changes.append((cell.coordinate,key))
            assert (cell.hyperlink.target if cell.hyperlink else None)==(other.hyperlink.target if other.hyperlink else None)
assert b['Papers']['L8'].value=='=AVERAGE(M8:N8)'
assert openpyxl.load_workbook(OUTPUT,data_only=True)['Papers']['L8'].value==7.6
with zipfile.ZipFile(SOURCE) as src,zipfile.ZipFile(OUTPUT) as dst:
    untouched=[name for name in src.namelist() if name not in [sheet_part,'xl/styles.xml']]
    assert all(src.read(name)==dst.read(name) for name in untouched)
report={'value_changes':value_changes,'style_changes':style_changes,'untouched_package_parts':len(untouched),'other_papers_unchanged':True,'average':7.6}
(ROOT/'workbook_verification.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report))
